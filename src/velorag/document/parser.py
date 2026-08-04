from __future__ import annotations

import io
import re
import uuid
from abc import ABC, abstractmethod
from pathlib import Path
from typing import Any

from velorag.core.types import Document


class BaseParser(ABC):
    @abstractmethod
    def parse(self, content: bytes, metadata: dict[str, Any] | None = None) -> Document:
        ...


class TextParser(BaseParser):
    def parse(self, content: bytes, metadata: dict[str, Any] | None = None) -> Document:
        text = content.decode("utf-8", errors="replace")
        return Document(
            id=uuid.uuid4().hex,
            content=text,
            metadata=metadata or {},
        )


class PDFParser(BaseParser):
    def parse(self, content: bytes, metadata: dict[str, Any] | None = None) -> Document:
        from PyPDF2 import PdfReader
        reader = PdfReader(io.BytesIO(content))
        text_parts = []
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
        return Document(
            id=uuid.uuid4().hex,
            content="\n\n".join(text_parts),
            metadata=metadata or {},
        )


class DocxParser(BaseParser):
    def parse(self, content: bytes, metadata: dict[str, Any] | None = None) -> Document:
        from docx import Document as DocxDoc
        doc = DocxDoc(io.BytesIO(content))
        text_parts = [p.text for p in doc.paragraphs if p.text.strip()]
        return Document(
            id=uuid.uuid4().hex,
            content="\n".join(text_parts),
            metadata=metadata or {},
        )


class MarkdownParser(BaseParser):
    def parse(self, content: bytes, metadata: dict[str, Any] | None = None) -> Document:
        text = content.decode("utf-8", errors="replace")
        cleaned = re.sub(r"!\[.*?\]\(.*?\)", "", text)
        cleaned = re.sub(r"\[([^\]]+)\]\(.*?\)", r"\1", cleaned)
        return Document(
            id=uuid.uuid4().hex,
            content=cleaned.strip(),
            metadata=metadata or {},
        )


PARSER_MAP: dict[str, type[BaseParser]] = {
    ".txt": TextParser,
    ".md": MarkdownParser,
    ".markdown": MarkdownParser,
    ".pdf": PDFParser,
    ".docx": DocxParser,
}
